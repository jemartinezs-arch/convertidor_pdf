import os
from docx import Document
from openpyxl import Workbook
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Preformatted


class ExportService:

    # ===============================
    # TEXTO → WORD (.docx)
    # ===============================
    def text_to_word(self, text, output_path):
        document = Document()
        document.add_paragraph(text)
        document.save(output_path)
        return output_path

    # ===============================
    # TEXTO → EXCEL (.xlsx)
    # Versión básica (divide por líneas)
    # ===============================
    def text_to_excel(self, text, output_path):
        wb = Workbook()
        ws = wb.active
        ws.title = "OCR Result"

        lines = text.split("\n")

        for row_index, line in enumerate(lines, start=1):
            ws.cell(row=row_index, column=1, value=line)

        wb.save(output_path)
        return output_path

    # ===============================
    # TEXTO → PDF EDITABLE
    # ===============================
    def text_to_pdf(self, text, output_path):
        doc = SimpleDocTemplate(output_path)
        elements = []

        styles = getSampleStyleSheet()
        style = styles["Normal"]

        lines = text.split("\n")

        for line in lines:
            elements.append(Paragraph(line, style))
            elements.append(Spacer(1, 0.2 * inch))

        doc.build(elements)
        return output_path

    # ===============================
    # GUARDAR TXT SIMPLE
    # ===============================
    def text_to_txt(self, text, output_path):
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(text)
        return output_path

    # ===============================
    # MÉTODO GENERAL EXPORTADOR
    # ===============================
    def export(self, text, format_type, output_folder="exports"):

        if not os.path.exists(output_folder):
            os.makedirs(output_folder)

        base_path = os.path.join(output_folder, "resultado")

        if format_type == "word":
            return self.text_to_word(text, base_path + ".docx")

        elif format_type == "excel":
            return self.text_to_excel(text, base_path + ".xlsx")

        elif format_type == "pdf":
            return self.text_to_pdf(text, base_path + ".pdf")

        elif format_type == "txt":
            return self.text_to_txt(text, base_path + ".txt")

        else:
            raise ValueError("Formato no soportado")